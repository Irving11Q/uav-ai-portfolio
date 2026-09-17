"""文档归一化层：把上传的文件转成 Markdown（**仅供索引使用**）。

## 为什么要做这一层

用户视角与机器视角是分离的：

- **用户视角**：上传、下载、在知识库列表里看到的，始终是原始文件（原样保存在
  `data/uploads/{id}_{filename}`，一个字节都不改）。用户看到的是他熟悉的 Word 样子。
- **机器视角**：系统内部把它转成 Markdown 作为「中间表示（IR）」，因为纯文本会丢掉
  标题层级和表格结构，导致后面的分块只能按字数硬切——表格被拦腰切断、续块没有表头。

Markdown 是这套系统里唯一同时满足两件事的格式：**对人可读、对切块器可结构化**。
所以它只出现在索引链路的中间，不出现在用户界面上。

## 转换内容

| Word 元素 | Markdown 输出 |
|---|---|
| Heading 1-6 / 标题 1-6 | `#` ~ `######` |
| 表格 | Markdown 表格（首行当表头） |
| 项目符号 / 编号列表 | `- ` / `1. ` |
| 加粗的行内文本 | `**粗体**` |
| 手动加粗的短行（无样式标题时） | `### 小标题`（启发式兜底） |

零新依赖，只用到已安装的 python-docx。
"""
import re

# Word 内置标题样式：英文 "Heading 1"，中文版 "标题 1"
_HEADING_RE = re.compile(r"^(?:heading|标题)\s*([1-6])$", re.IGNORECASE)
# 手动加粗短行的启发式判定：多长算短、以及哪些结尾不算标题
_BOLD_HEADING_MAX_LEN = 24
_BOLD_HEADING_END_PUNCT = "。！？!?；;，,"

# 列表样式识别（英文小写去空格后比对）
_LIST_STYLE_PREFIXES = ("listbullet", "listnumber")
_LIST_STYLE_NAMES = ("列表项", "项目符号", "编号列表", "符号列表")
_LIST_PARAGRAPH_NAMES = ("listparagraph", "列表段落")


def rows_to_markdown(rows: list[list[str]]) -> str:
    """把二维数据转成 Markdown 表格（首行作表头）。csv / xlsx / docx 表格共用。"""
    clean = [[_escape_cell(c) for c in row] for row in rows]
    clean = [r for r in clean if any(c for c in r)]
    if not clean:
        return ""
    width = max(len(r) for r in clean)
    clean = [r + [""] * (width - len(r)) for r in clean]
    lines = ["| " + " | ".join(clean[0]) + " |", "|" + "---|" * width]
    for row in clean[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def docx_to_markdown(path: str) -> str:
    """把 .docx 转成 Markdown，保留标题层级 / 表格 / 列表 / 加粗。"""
    try:
        from docx import Document as DocxDocument
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as e:  # pragma: no cover
        raise ValueError("缺少依赖 python-docx，无法解析 docx") from e

    doc = DocxDocument(path)

    # 按 body 的真实顺序遍历「段落 + 表格」。
    # 注意：不能用 doc.paragraphs + doc.tables 分开取——那会把所有表格堆到文档末尾，
    # 表格与它上文的位置关系就丢了，检索时表格会脱离上下文。
    blocks = []
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            blocks.append(("p", Paragraph(child, doc)))
        elif isinstance(child, CT_Tbl):
            blocks.append(("t", Table(child, doc)))

    # 整篇若用了样式标题，就以样式为准；否则对「整行加粗的短行」启用启发式，
    # 兼容那种手动加粗当标题、没用样式的手写 Word。
    has_styled_heading = any(
        kind == "p" and _heading_level(_style_name(obj)) for kind, obj in blocks
    )

    parts: list[tuple[str, bool]] = []  # (块文本, 是否列表项)
    for kind, obj in blocks:
        if kind == "t":
            md = _table_to_markdown(obj)
            if md:
                parts.append((md, False))
            continue
        line = _paragraph_to_markdown(obj, bold_heading_fallback=not has_styled_heading)
        if line:
            parts.append((line, _is_list_line(line)))

    # 连续列表项之间不留空行——留了会被 Markdown 解析成多个独立列表，渲染出来是断开的
    text = ""
    for i, (line, is_list) in enumerate(parts):
        if i == 0:
            text = line
            continue
        tight = is_list and parts[i - 1][1]
        text += ("\n" if tight else "\n\n") + line

    # 折叠连续空行，保证段落边界干净（空行是切块时的首选分隔符）
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _is_list_line(line: str) -> bool:
    """生成的 Markdown 行是否列表项（看我们自己的前缀标记）。"""
    if line.startswith("- "):
        return True
    return len(line) > 3 and line[0].isdigit() and line[1:3] == ". "


# ---------------------------------------------------------------- 内部实现


def _escape_cell(value) -> str:
    """单元格文本压平为单行并转义竖线，避免破坏 Markdown 表格列结构。"""
    txt = " ".join(str(value or "").split())
    return txt.replace("|", r"\|")


def _style_name(p) -> str:
    """取段落样式名，取不到时返回空串（异常不向外抛）。"""
    try:
        return (p.style.name or "").strip()
    except Exception:
        return ""


def _heading_level(style_name: str) -> int | None:
    """样式名 → 标题层级；不是标题样式返回 None。"""
    m = _HEADING_RE.match(style_name or "")
    if m:
        return int(m.group(1))
    if (style_name or "").strip().lower() in ("title", "标题"):
        return 1
    return None


def _is_list_item(p) -> bool:
    """段落是否属于列表。

    两步判断，缺一不可：
    1. 自动编号的段落会在 pPr 里写 numPr —— 这是最可靠的信号；
    2. 但 Word 内置的 List Bullet / List Number **样式**不会在段落里写 numPr，
       编号定义在 styles.xml 中，所以必须再看样式名兜底（实测 List Bullet 只能靠这条）。
    "List Paragraph / 列表段落" 只是缩进用的普通段落，不算列表项。
    """
    try:
        pPr = p._p.pPr
        if pPr is not None and pPr.numPr is not None:
            return True
    except Exception:
        pass
    name = _style_name(p)
    lowered = name.lower().replace(" ", "")
    if lowered in _LIST_PARAGRAPH_NAMES or name in _LIST_PARAGRAPH_NAMES:
        return False
    return lowered.startswith(_LIST_STYLE_PREFIXES) or name in _LIST_STYLE_NAMES


def _is_ordered_list(p) -> bool:
    """有序列号用 `1. `，否则用 `- `。"""
    s = _style_name(p).lower().replace(" ", "")
    return "number" in s or "编号" in _style_name(p)


def _run_text(run) -> str:
    """把单个 run 转成 Markdown 片段：加粗包 **，保留首尾空格。"""
    t = run.text
    if not t:
        return ""
    core = t.strip()
    if not core:
        return t
    if run.bold:
        lead = t[: len(t) - len(t.lstrip())]
        trail = t[len(t.rstrip()) :]
        return f"{lead}**{core}**{trail}"
    return t


def _paragraph_text(p) -> str:
    """段落转 Markdown 行内文本；runs 取不到内容时退回 p.text（如超链接）。"""
    text = "".join(_run_text(r) for r in p.runs).replace("\n", " ")
    if not text.strip():
        text = (p.text or "").replace("\n", " ")
    return text.strip()


def _paragraph_to_markdown(p, bold_heading_fallback: bool = False) -> str:
    """单个段落 → Markdown 行。"""
    text = _paragraph_text(p)
    if not text:
        return ""
    style = _style_name(p)

    level = _heading_level(style)
    if level:
        return "#" * level + " " + _strip_wrapping_bold(text).lstrip("#").strip()

    if _is_list_item(p):
        return ("1. " if _is_ordered_list(p) else "- ") + text

    if bold_heading_fallback and _looks_like_bold_heading(p, text):
        # 整行加粗的行升格为标题时，要去掉外层 ** —— 否则会输出「### **标题**」这种冗余标记
        return "### " + _strip_wrapping_bold(text)

    return text


def _strip_wrapping_bold(text: str) -> str:
    """剥掉「整行被一对 ** 包住」的加粗标记，保留行内其它加粗。"""
    m = re.fullmatch(r"\*\*(.+)\*\*", text.strip(), re.S)
    return m.group(1) if m else text.strip()


def _looks_like_bold_heading(p, text: str) -> bool:
    """启发式：整行加粗、够短、不以句末标点收尾 → 当作小标题。"""
    if len(text) > _BOLD_HEADING_MAX_LEN or text.endswith(tuple(_BOLD_HEADING_END_PUNCT)):
        return False
    runs = [r for r in p.runs if r.text.strip()]
    return bool(runs) and all(r.bold for r in runs)


def _table_to_markdown(table) -> str:
    """Word 表格 → Markdown 表格；横向合并的单元格只取一次，避免内容重复。"""
    rows: list[list[str]] = []
    for row in table.rows:
        cells, last_tc = [], None
        for cell in row.cells:
            tc = getattr(cell, "_tc", None)
            if tc is not None and last_tc is not None and tc is last_tc:
                continue  # 横向合并单元格会重复返回同一个 tc，只取首个
            last_tc = tc
            cells.append(_escape_cell(cell.text))
        if any(cells):
            rows.append(cells)
    return rows_to_markdown(rows)
